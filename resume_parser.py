
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Extract text from an uploaded PDF or DOCX file."""
    extension = Path(filename).suffix.lower()

    if extension not in {".pdf", ".docx"}:
        raise ValueError("Please upload a PDF or DOCX file.")

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError("The file must be no larger than 5 MB.")

    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(file_bytes))

            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported.")

            pages = []
            for page in reader.pages:
                pages.append(page.extract_text() or "")

            text = "\n".join(pages)

        else:
            document = Document(BytesIO(file_bytes))
            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            # Also extract text from tables, which often contain resume details.
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(
                        " ".join(cell.text for cell in row.cells)
                    )

            text = "\n".join(paragraphs)

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Could not read this file. Please check that it is a valid PDF or DOCX."
        ) from exc

    if not text.strip():
        raise ValueError(
            "No text was extracted. If this is a scanned PDF, OCR is required."
        )

    return text
